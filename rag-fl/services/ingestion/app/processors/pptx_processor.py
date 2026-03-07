"""
services/ingestion/app/processors/pptx_processor.py
Handles PPTX by:
1. Extracting slide count from pptx metadata (cheap, no conversion)
2. Converting to PDF via Gotenberg LibreOffice service
3. Delegating PDF processing to PDFProcessor for page count verification
"""
import io
import os
import requests
from pptx import Presentation
from .base import BaseProcessor, NormalizedOutput
from .pdf_processor import PDFProcessor


def _convert_to_pdf(file_bytes: bytes, filename: str) -> bytes:
    libreoffice_url = os.getenv("LIBREOFFICE_URL", "http://libreoffice:3000")
    response = requests.post(
        f"{libreoffice_url}/forms/libreoffice/convert",
        files={"files": (filename, file_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        timeout=120,
    )
    response.raise_for_status()
    return response.content


class PPTXProcessor(BaseProcessor):

    def validate(self, file_bytes: bytes) -> bool:
        if not file_bytes:
            return False
        try:
            prs = Presentation(io.BytesIO(file_bytes))
            return len(prs.slides) > 0
        except Exception:
            return False

    def normalize(self, file_bytes: bytes, filename: str) -> NormalizedOutput:
        prs = Presentation(io.BytesIO(file_bytes))
        slide_count = len(prs.slides)

        pdf_bytes = _convert_to_pdf(file_bytes, filename)
        pdf_proc = PDFProcessor()
        pdf_normalized = pdf_proc.normalize(pdf_bytes, filename.replace(".pptx", ".pdf"))

        return NormalizedOutput(
            file_bytes=pdf_bytes,
            mime_type="application/pdf",
            page_count=pdf_normalized.page_count,
            metadata={"original_slide_count": slide_count},
            provenance={
                "original_format": "pptx",
                "slide_count": slide_count,
                "converted_via": "gotenberg",
                "converted_page_count": pdf_normalized.page_count,
            },
        )

    def extract_provenance(self, normalized: NormalizedOutput) -> dict:
        return normalized.provenance
